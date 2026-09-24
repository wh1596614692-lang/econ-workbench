import io
import json
import os
import subprocess
import sys
import time
import zipfile
from pathlib import Path
import pandas as pd
from fastapi.testclient import TestClient
from econworkbench.api import app

ROOT=Path(__file__).resolve().parents[1]


def wait_for(c,jid):
    for _ in range(240):
        result=c.get('/api/jobs/'+jid).json()
        if result['status']!='running':
            assert result['status']=='completed',result
            return result['result']
        time.sleep(.15)
    raise AssertionError('job timed out')


def test_upload_to_export_and_isolation(tmp_path):
    with TestClient(app) as c:
        token=c.post('/api/session',json={}).json()['token']
        c.headers.update({'Authorization':'Bearer '+token})
        data=(ROOT/'examples'/'synthetic_cross.csv').read_bytes()
        ds=c.post('/api/upload',files={'file':('study.csv',data,'text/csv')}).json()
        assert ds['rows']==320
        transformed=c.post(f'/api/datasets/{ds["id"]}/transform',json={'transform':{'op':'center','columns':['education'],'target':'centered_education'}}).json()
        assert transformed['log'][0]['affected']>0
        cfg={'name':'端到端验收','y':'log_income','x':['centered_education'],'controls':['experience'],'se':'HC3','group':'region'}
        response=c.post('/api/runs',json={'dataset_id':transformed['id'],'config':cfg})
        assert response.status_code==200,response.text
        result=wait_for(c,response.json()['job_id'])
        assert result['nobs']==317
        assert any(t['status']=='已完成' for t in result['tests'])
        z=c.post('/api/export',json={'run_ids':[result['id']]})
        assert z.status_code==200,z.text[:200] if z.status_code!=200 else ''
        with zipfile.ZipFile(io.BytesIO(z.content)) as archive:
            assert {'report.html','results.xlsx','reproduce.py','regression_table.tex'}.issubset(archive.namelist())
            saved=json.loads(archive.read('results/model_1.json'))
            assert saved['coefficients']==result['coefficients']
            excel=pd.read_excel(io.BytesIO(archive.read('results.xlsx')),sheet_name='1_系数')
            assert abs(excel.iloc[1].coefficient-result['coefficients'][1]['coefficient'])<1e-12
            archive.extractall(tmp_path)
        env=dict(os.environ,PYTHONIOENCODING='utf-8')
        run=subprocess.run([sys.executable,str(tmp_path/'reproduce.py')],cwd=tmp_path,env=env,capture_output=True,text=True,timeout=90,encoding='utf-8')
        assert run.returncode==0,run.stderr
        assert 'OK' in run.stdout
        another=c.post('/api/session',json={}).json()['token']
        response=c.post(f'/api/datasets/{ds["id"]}/inspect',json={},headers={'Authorization':'Bearer '+another})
        assert response.status_code==404
        c.delete('/api/session',headers={'Authorization':'Bearer '+another})
        c.delete('/api/session')
        assert c.get('/api/runs').status_code==401


def test_cancel_and_excel_sheets():
    with TestClient(app) as c:
        token=c.post('/api/session',json={}).json()['token'];c.headers.update({'Authorization':'Bearer '+token})
        data=(ROOT/'examples'/'synthetic_workbook.xlsx').read_bytes()
        ds=c.post('/api/upload',files={'file':('data.xlsx',data)},data={'sheet':'面板_合成演示数据'}).json()
        assert ds['rows']==1200 and ds['sheet']=='面板_合成演示数据'
        cfg=json.loads((ROOT/'examples'/'configs.json').read_text(encoding='utf-8'))['gmm']
        job=c.post('/api/runs',json={'dataset_id':ds['id'],'config':cfg}).json()['job_id']
        cancelled=c.delete('/api/jobs/'+job)
        assert cancelled.json()['status']=='cancelled'
        assert c.get('/api/jobs/'+job).json()['status']=='cancelled'
        c.delete('/api/session')


def test_csv_formula_and_html_escaping():
    from econworkbench.exports import safe_frame
    from openpyxl import load_workbook
    x=safe_frame(pd.DataFrame({'x':['=HYPERLINK("evil")','+SUM(A1)','@X','-1+2']}))
    assert all(v.startswith("'") for v in x.x)
    b=io.BytesIO();x.to_excel(b,index=False)
    workbook=load_workbook(io.BytesIO(b.getvalue()))
    assert all(cell.data_type!='f' for row in workbook.active.iter_rows() for cell in row)


def test_common_sample_reestimation():
    """Different complete-case samples must become the explicit intersection."""
    with TestClient(app) as c:
        token=c.post('/api/session',json={}).json()['token']
        c.headers.update({'Authorization':'Bearer '+token})
        frame=pd.read_csv(ROOT/'examples'/'synthetic_cross.csv')
        frame.loc[:19,'experience']=float('nan')
        ds=c.post('/api/upload',files={'file':('comparison.csv',frame.to_csv(index=False).encode())}).json()
        runs=[]
        for controls in [[],['experience']]:
            job=c.post('/api/runs',json={'dataset_id':ds['id'],'config':{'y':'log_income','x':['education'],'controls':controls,'se':'HC3'}}).json()['job_id']
            runs.append(wait_for(c,job))
        expected=sorted(set(runs[0]['sample_ids']) & set(runs[1]['sample_ids']))
        assert runs[0]['nobs']>runs[1]['nobs']
        specs=c.post('/api/common-sample',json={'run_ids':[r['id'] for r in runs]}).json()
        assert specs['nobs']==len(expected)
        for spec in specs['specifications']:
            job=c.post('/api/runs',json=spec).json()['job_id']
            result=wait_for(c,job)
            assert result['sample_ids']==expected
            assert result['nobs']==len(expected)
        assert c.get('/docs').status_code==200
        assert '/api/runs' in c.get('/docs').text
        c.delete('/api/session')
